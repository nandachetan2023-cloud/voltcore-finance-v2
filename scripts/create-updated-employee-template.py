#!/usr/bin/env python3
"""
Create updated employee import template with new salary fields
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from copy import copy

# Read the existing template
wb = openpyxl.load_workbook('excels/employees.xlsx')
ws = wb.active

# Get existing headers
existing_headers = [cell.value for cell in ws[1]]

# Define new columns to add after Employee ID (column 2)
new_columns = [
    ('Token Number', 'Unique token number (different from Employee ID)'),
    ('Workmen Sl. No.', 'Workmen serial number for salary sheet'),
    ('Monthly Gross Salary', 'Monthly gross salary amount'),
    ('Nature of Designation', 'Skilled/Unskilled/Semi-skilled'),
]

# Insert new columns after column B (Employee ID)
insert_position = 3  # After Employee ID (column B)

for i, (col_name, description) in enumerate(new_columns):
    ws.insert_cols(insert_position + i)
    
    # Copy formatting from adjacent column
    source_col = ws.cell(1, 2)  # Employee ID column
    target_col = ws.cell(1, insert_position + i)
    
    # Set header
    target_col.value = col_name
    
    # Copy formatting
    if source_col.has_style:
        target_col.font = copy(source_col.font)
        target_col.fill = copy(source_col.fill)
        target_col.alignment = copy(source_col.alignment)
        target_col.border = copy(source_col.border)
    
    # Add comment/note
    target_col.comment = openpyxl.comments.Comment(description, 'System')

# Update column widths
ws.column_dimensions[openpyxl.utils.get_column_letter(3)].width = 15  # Token Number
ws.column_dimensions[openpyxl.utils.get_column_letter(4)].width = 18  # Workmen Sl. No.
ws.column_dimensions[openpyxl.utils.get_column_letter(5)].width = 20  # Monthly Gross Salary
ws.column_dimensions[openpyxl.utils.get_column_letter(6)].width = 22  # Nature of Designation

# Add sample data in row 2 (if it exists)
if ws.max_row >= 2:
    ws.cell(2, 3).value = 'TKN001'  # Token Number
    ws.cell(2, 4).value = 'WM001'   # Workmen Sl. No.
    ws.cell(2, 5).value = 15000     # Monthly Gross Salary
    ws.cell(2, 6).value = 'Skilled' # Nature of Designation

# Save the updated template
wb.save('excels/employees_updated.xlsx')
print('✅ Created updated employee template: excels/employees_updated.xlsx')

# Print column mapping
print('\n📋 Updated Column Structure:')
print('-' * 80)
for i in range(1, ws.max_column + 1):
    col_letter = openpyxl.utils.get_column_letter(i)
    col_name = ws.cell(1, i).value
    if col_name:
        marker = '🆕' if i in [3, 4, 5, 6] else '  '
        print(f'{marker} {col_letter:3} | {col_name}')

print('\n' + '=' * 80)
print('NEW COLUMNS ADDED:')
print('=' * 80)
print('C  | Token Number          - Unique token (different from Employee ID)')
print('D  | Workmen Sl. No.       - Workmen serial number')
print('E  | Monthly Gross Salary  - Monthly gross salary amount')
print('F  | Nature of Designation - Skilled/Unskilled/Semi-skilled')
print('\nNote: Sl.No. (Column A) will be auto-generated during import')
