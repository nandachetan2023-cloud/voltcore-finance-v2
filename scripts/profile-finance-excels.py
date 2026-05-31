from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl


def _norm_header(value: Any) -> str:
    if value is None:
        return ""
    s = str(value).strip()
    s = re.sub(r"\s+", " ", s)
    return s


def _is_empty_cell(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and value.strip() == "":
        return True
    return False


@dataclass(frozen=True)
class HeaderGuess:
    row_idx: int
    score: float
    nonempty: int
    max_col: int
    headers: List[str]


def _guess_header_row(ws, max_scan_rows: int = 40, max_scan_cols: int = 400) -> Optional[HeaderGuess]:
    max_row = int(ws.max_row or 0)
    max_col = int(ws.max_column or 0)
    if max_row <= 0 or max_col <= 0:
        return None

    scan_rows = min(max_row, max_scan_rows)
    scan_cols = min(max_col, max_scan_cols)

    best: Optional[HeaderGuess] = None

    # Use iter_rows for speed (cell() access is slow in openpyxl)
    rows_iter = ws.iter_rows(min_row=1, max_row=scan_rows, min_col=1, max_col=scan_cols, values_only=True)
    for r, row_vals in enumerate(rows_iter, start=1):
        row_vals = list(row_vals)
        headers = [_norm_header(v) for v in row_vals]
        nonempty = sum(1 for h in headers if h != "")

        if nonempty < 2:
            continue

        # Prefer rows that look like a header:
        # - many non-empty cells
        # - lower row index (earlier)
        # - more "stringy" cells
        strish = 0
        for v in row_vals:
            if isinstance(v, str):
                strish += 1
            elif v is not None and not isinstance(v, (int, float)):
                # dates/datetimes count as header-like less than strings but more than numbers
                strish += 0.5

        # Penalize rows that look like data (mostly numbers)
        numericish = 0
        for v in row_vals:
            if isinstance(v, (int, float)) and v is not None:
                numericish += 1

        # Uniqueness helps (headers often unique)
        uniq = len({h.lower() for h in headers if h})

        score = (
            nonempty * 2.0
            + uniq * 0.5
            + strish * 0.2
            - numericish * 0.6
            - (r - 1) * 0.15
        )

        guess = HeaderGuess(row_idx=r, score=score, nonempty=nonempty, max_col=scan_cols, headers=headers)
        if best is None or guess.score > best.score:
            best = guess

    return best


def _dedupe_headers(headers: List[str]) -> List[str]:
    seen: Dict[str, int] = {}
    out: List[str] = []
    for i, h in enumerate(headers, start=1):
        base = h.strip()
        if base == "":
            base = f"UNNAMED_{i}"
        key = base.lower()
        seen[key] = seen.get(key, 0) + 1
        if seen[key] > 1:
            out.append(f"{base} ({seen[key]})")
        else:
            out.append(base)
    return out


def _effective_max_col(ws, max_scan_rows: int = 60) -> int:
    """
    openpyxl's ws.max_column can be huge due to formatting; derive a practical used column count
    by scanning the first N rows and taking the last non-empty cell index.

    This aims to capture ALL actual data columns without iterating 16k columns per sheet.
    """
    max_row = int(ws.max_row or 0)
    if max_row <= 0:
        return 0

    scan_rows = min(max_row, max_scan_rows)
    max_col_seen = 0
    for row in ws.iter_rows(min_row=1, max_row=scan_rows, values_only=True):
        # Trim trailing empties
        last = 0
        for idx in range(len(row), 0, -1):
            if not _is_empty_cell(row[idx - 1]):
                last = idx
                break
        if last > max_col_seen:
            max_col_seen = last

    # Fallback to at least 1 if sheet has structure but empties in the first rows
    if max_col_seen == 0:
        max_col_seen = min(int(ws.max_column or 0), 50)

    return max_col_seen


def _sheet_preview_rows(ws, start_row: int, rows: int, cols: int) -> List[List[str]]:
    max_row = int(ws.max_row or 0)
    max_col = int(ws.max_column or 0)
    end_row = min(max_row, start_row + rows - 1)
    cols = min(max_col, cols)

    out: List[List[str]] = []
    for r in range(start_row, end_row + 1):
        row = []
        for c in range(1, cols + 1):
            v = ws.cell(r, c).value
            if v is None:
                row.append("")
            else:
                s = str(v)
                s = s.replace("\r\n", "\n")
                row.append(s)
        out.append(row)
    return out


def profile_workbook(path: Path) -> Dict[str, Any]:
    info: Dict[str, Any] = {
        "file": path.name,
        "path": str(path),
        "size_bytes": path.stat().st_size,
        "sheets": [],
    }

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        info["error"] = f"{type(e).__name__}: {e}"
        return info

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        max_row = int(ws.max_row or 0)
        eff_col = _effective_max_col(ws)
        max_col = int(ws.max_column or 0)

        guess = _guess_header_row(ws, max_scan_cols=min(eff_col or 0, 600) or 200)
        header_row = guess.row_idx if guess else 1

        # Capture ALL columns up to effective max col (derived from actual non-empty cells)
        # Using iter_rows avoids slow per-cell access.
        eff_col = max(eff_col, 0)
        if eff_col <= 0:
            raw_headers = []
        else:
            header_vals = next(
                ws.iter_rows(min_row=header_row, max_row=header_row, min_col=1, max_col=eff_col, values_only=True),
                (),
            )
            raw_headers = [_norm_header(v) for v in list(header_vals)]
        headers = _dedupe_headers(raw_headers)

        # Also capture the top rows so multi-row headers can be reconstructed manually if needed
        top_rows = _sheet_preview_rows(ws, start_row=1, rows=min(8, max_row if max_row else 8), cols=min(eff_col or 0, 60))

        info["sheets"].append(
            {
                "name": sheet_name,
                "max_row": max_row,
                "max_col": max_col,
                "effective_max_col": eff_col,
                "header_row_guess": header_row,
                "header_score": round(float(guess.score), 3) if guess else None,
                "columns": headers,
                "columns_count": len(headers),
                "top_rows_preview": top_rows,
            }
        )

    return info


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True, help="Folder containing finance excels")
    ap.add_argument("--out-json", required=True, help="Output JSON path")
    ap.add_argument("--out-md", required=True, help="Output Markdown path")
    args = ap.parse_args()

    base = Path(args.input)
    if not base.exists():
        raise SystemExit(f"Not found: {base}")

    files = sorted([p for p in base.rglob("*") if p.suffix.lower() in {".xlsx", ".xlsm"}])

    results = [profile_workbook(p) for p in files]

    out_json = Path(args.out_json)
    out_json.parent.mkdir(parents=True, exist_ok=True)
    out_json.write_text(json.dumps(results, indent=2, ensure_ascii=False), encoding="utf-8")

    # Markdown summary (compact but complete list of columns)
    lines: List[str] = []
    lines.append("# Finance Excel Profile")
    lines.append("")
    lines.append(f"Input folder: `{base}`")
    lines.append(f"Workbooks found: **{len(results)}**")
    lines.append("")

    for wb in results:
        lines.append(f"## {wb.get('file')}")
        if wb.get("error"):
            lines.append(f"- Error: `{wb['error']}`")
            lines.append("")
            continue
        lines.append(f"- Path: `{wb.get('path')}`")
        lines.append(f"- Size: {wb.get('size_bytes')} bytes")
        for sh in wb.get("sheets", []):
            lines.append(f"### Sheet: {sh.get('name')}")
            lines.append(f"- Rows: {sh.get('max_row')}, Cols: {sh.get('max_col')}")
            lines.append(f"- Header row guess: {sh.get('header_row_guess')} (score {sh.get('header_score')})")
            lines.append(f"- Columns ({sh.get('columns_count')}):")
            for col in sh.get("columns", []):
                lines.append(f"  - {col}")
            lines.append("")

    out_md = Path(args.out_md)
    out_md.parent.mkdir(parents=True, exist_ok=True)
    out_md.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    main()
