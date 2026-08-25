#!/usr/bin/env python3
"""
Generate expense import template for pritishon ERP
This script creates an Excel file with all required columns for HO personal expenses
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def create_expense_template():
    """Create the expense import template Excel file"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HO Expenditure Data"

    # Title and description
    ws.cell(row=1, column=1, value="HO Personal Expenses Import Template")
    ws.cell(row=2, column=1, value="Please fill in the following columns:")
    ws.cell(row=2, column=1).font = Font(bold=True, size=12)
    ws.cell(row=3, column=1, value="")

    # Headers
    headers = [
        "Date",
        "SiteType (HO)",
        "Category",
        "ItemName",
        "Description",
        "TotalAmount",
        "ReceivedAmount",
        "GSTAmount",
        "TDSAmount",
        "BillNo",
        "ApprovalStatus"
    ]

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num, value=header)
        cell.font = Font(bold=True, size=11)
        cell.fill = PatternFill(start_color="f5a623", end_color="e8891a", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Note section
    ws.cell(row=6, column=1, value="Notes:")
    ws.cell(row=7, column=1, value="1. Date format: dd-mm-yyyy (e.g., 12-07-2026)")
    ws.cell(row=8, column=1, value="2. SiteType: Must be 'HO'")
    ws.cell(row=9, column=1, value="3. Category: GAP FOODING, ADVANCE, ABF, ETC")
    ws.cell(row=10, column=1, value="4. TotalAmount: Total expense amount")
    ws.cell(row=11, column=1, value="5. ReceivedAmount: Amount received from HO")
    ws.cell(row=12, column=1, value="6. GSTAmount: GST amount(Tax)")
    ws.cell(row=13, column=1, value="7. TDSAmount: TDS amount")
    ws.cell(row=14, column=1, value="8. BillNo: Bill or invoice number")
    ws.cell(row=15, column=1, value="9. ApprovalStatus: Draft, Pending, Approved, Rejected, Paid")

    for col_num, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(col_num)].width = 18

    # Adjust row heights
    ws.row_dimensions[4].height = 25
    for i in range(6, 16):
        ws.row_dimensions[i].height = 20

    # Save the file
    wb.save("expense-import-template.xlsx")
    print("✓ Template created successfully: expense-import-template.xlsx")

if __name__ == "__main__":
    try:
        create_expense_template()
    except ImportError:
        print("Error: openpyxl library not found")
        print("Please install it with: pip install openpyxl")
    except Exception as e:
        print(f"Error creating template: {e}")